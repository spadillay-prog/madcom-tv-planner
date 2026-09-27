import calendar
from datetime import date, datetime, timedelta, time
import io
import math
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="MADCOM — Planificador TV", page_icon="📺", layout="wide"
)

# -------------------------------------------------------------
# 1. MOTOR DE EXTRACCIÓN DE TARIFARIOS Y RATINGS POR CLIENTE
# -------------------------------------------------------------
@st.cache_data
def load_compara_inventory():
    file_path = "Valor TV Abierta Compara.xlsx"
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_progs = []

    # CHV
    ws_chv = wb["CHV"]
    dur_map_chv = {
        70: 9, 60: 12, 55: 13, 50: 14, 45: 15, 40: 16,
        35: 17, 30: 18, 25: 19, 20: 20, 15: 21, 10: 22, 5: 23,
    }
    for r in range(10, ws_chv.max_row + 1):
        prog = ws_chv.cell(r, 3).value
        if not prog:
            continue
        bloque = str(ws_chv.cell(r, 2).value or "Off").strip()
        dias = str(ws_chv.cell(r, 4).value or "L-V").strip()
        hra_val = ws_chv.cell(r, 5).value
        hra = (
            hra_val.strftime("%H:%M")
            if isinstance(hra_val, (datetime, time))
            else str(hra_val or "")
        )
        rat = ws_chv.cell(r, 11).value or 0.0
        try:
            rat = float(rat)
        except:
            rat = 0.0

        tariffs = {}
        for sec, c_idx in dur_map_chv.items():
            val = ws_chv.cell(r, c_idx).value
            try:
                tariffs[sec] = float(val) if val else 0.0
            except:
                tariffs[sec] = 0.0

        all_progs.append(
            {
                "Canal": "CHV",
                "Programa": str(prog).strip(),
                "Bloque": "PRIME" if "prime" in bloque.lower() else "OFF PRIME",
                "Días": dias,
                "Horario": hra,
                "Rating": rat,
                "Tarifas": tariffs,
            }
        )

    # Mega
    ws_mega = wb["Mega"]
    dur_map_mega = {}
    for c in range(7, 25):
        val = ws_mega.cell(10, c).value
        if val and isinstance(val, (int, float)):
            dur_map_mega[int(val)] = c

    for r in range(12, ws_mega.max_row + 1):
        prog = ws_mega.cell(r, 2).value
        if not prog:
            continue
        dias = str(ws_mega.cell(r, 3).value or "L-V").strip()
        hra = str(ws_mega.cell(r, 4).value or "").strip()
        bloque = str(ws_mega.cell(r, 5).value or "Off").strip()
        rat = ws_mega.cell(r, 6).value or 0.0
        try:
            rat = float(rat)
        except:
            rat = 0.0

        tariffs = {}
        for sec, c_idx in dur_map_mega.items():
            val = ws_mega.cell(r, c_idx).value
            try:
                tariffs[sec] = float(val) if val else 0.0
            except:
                tariffs[sec] = 0.0

        all_progs.append(
            {
                "Canal": "Mega",
                "Programa": str(prog).strip(),
                "Bloque": "PRIME" if "prime" in bloque.lower() else "OFF PRIME",
                "Días": dias,
                "Horario": hra,
                "Rating": rat,
                "Tarifas": tariffs,
            }
        )

    # TVN
    ws_tvn = wb["TVN"]
    dur_map_tvn = {}
    for c in range(14, 28):
        h_val = str(ws_tvn.cell(6, c).value or "")
        digits = "".join([ch for ch in h_val if ch.isdigit()])
        if digits:
            dur_map_tvn[int(digits)] = c

    for r in range(7, ws_tvn.max_row + 1):
        prog = ws_tvn.cell(r, 2).value
        if not prog:
            continue
        dias = str(ws_tvn.cell(r, 3).value or "L-V").strip()
        ini_val = ws_tvn.cell(r, 5).value
        fin_val = ws_tvn.cell(r, 6).value
        ini = (
            ini_val.strftime("%H:%M")
            if isinstance(ini_val, (datetime, time))
            else str(ini_val or "")
        )
        fin = (
            fin_val.strftime("%H:%M")
            if isinstance(fin_val, (datetime, time))
            else str(fin_val or "")
        )
        hra = f"{ini} - {fin}" if ini and fin else ini
        bloque = str(ws_tvn.cell(r, 12).value or "Off").strip()
        rat = ws_tvn.cell(r, 13).value or 0.0
        try:
            rat = float(rat)
        except:
            rat = 0.0

        tariffs = {}
        for sec, c_idx in dur_map_tvn.items():
            val = ws_tvn.cell(r, c_idx).value
            try:
                tariffs[sec] = float(val) if val else 0.0
            except:
                tariffs[sec] = 0.0

        all_progs.append(
            {
                "Canal": "TVN",
                "Programa": str(prog).strip(),
                "Bloque": "PRIME" if "prime" in bloque.lower() else "OFF PRIME",
                "Días": dias,
                "Horario": hra,
                "Rating": rat,
                "Tarifas": tariffs,
            }
        )

    # Canal 13
    ws_c13 = wb["Canal 13"]
    dur_map_c13 = {}
    for c in range(6, 25):
        h_val = str(ws_c13.cell(2, c).value or "")
        digits = "".join([ch for ch in h_val if ch.isdigit()])
        if digits:
            dur_map_c13[int(digits)] = c

    for r in range(3, ws_c13.max_row + 1):
        prog = ws_c13.cell(r, 2).value
        if not prog or "PROGRAMAS" in str(prog).upper():
            continue
        dias = str(ws_c13.cell(r, 3).value or "L-V").strip()
        bloque = str(ws_c13.cell(r, 4).value or "Off").strip()
        rat = ws_c13.cell(r, 5).value or 0.0
        try:
            rat = float(rat)
        except:
            rat = 0.0

        tariffs = {}
        for sec, c_idx in dur_map_c13.items():
            val = ws_c13.cell(r, c_idx).value
            try:
                tariffs[sec] = float(val) if val else 0.0
            except:
                tariffs[sec] = 0.0

        all_progs.append(
            {
                "Canal": "Canal 13",
                "Programa": str(prog).strip(),
                "Bloque": "PRIME" if "prime" in bloque.lower() else "OFF PRIME",
                "Días": dias,
                "Horario": "—",
                "Rating": rat,
                "Tarifas": tariffs,
            }
        )

    return all_progs


progs_db = load_compara_inventory()

# -------------------------------------------------------------
# 2. INTERFAZ: PARÁMETROS DE CAMPAÑA
# -------------------------------------------------------------
st.title("📺 MADCOM — Planificador & Cotizador TV Abierta")

with st.sidebar:
    st.header("🏢 Selección de Cuenta")
    cliente_sel = st.selectbox(
        "Cliente Activo",
        options=["Compara Online", "+ Agregar nuevo cliente..."],
        index=0,
    )
    campana_nombre = st.text_input("Campaña:", value="Pauta Mes Tipo 2026")
    version_pauta = st.text_input("Versión:", value="V1.0")

    st.markdown("---")
    st.header("📅 Período de Campaña")
    default_start = date(2026, 10, 1)
    default_end = date(2026, 10, 20)
    date_range = st.date_input(
        "Rango de Fechas:",
        value=(default_start, default_end),
        min_value=date(2026, 1, 1),
        max_value=date(2027, 12, 31),
    )

    if isinstance(date_range, tuple) and len(date_range) == 2:
        start_date, end_date = date_range
    else:
        start_date = date_range[0] if isinstance(date_range, tuple) else date_range
        end_date = start_date + timedelta(days=19)

    num_dias_campana = (end_date - start_date).days + 1
    st.caption(f"Duración: **{num_dias_campana} días de exhibición**.")

    st.markdown("---")
    st.header("💰 Objetivos de Medios")
    presupuesto_total = st.number_input(
        "Inversión Total Meta (CLP):",
        min_value=1000000,
        value=40000000,
        step=1000000,
        format="%d",
    )
    objetivo_trps = st.number_input(
        "Objetivo TRPs:", min_value=10, value=200, step=10
    )
    sov_prime_target = (
        st.slider("SOV TRPs Prime (%):", min_value=10, max_value=90, value=70) / 100.0
    )

    st.markdown("---")
    st.header("⏱️ Mix de Segundajes")
    AVAILABLE_SECS = [5, 10, 15, 20, 25, 30, 35, 40, 45, 50, 55, 60, 65, 70]
    multi_segundaje = st.checkbox("Asignar segundajes por bloque", value=True)

    if multi_segundaje:
        seg_prime = st.selectbox(
            "Segundaje en Bloque PRIME:", options=AVAILABLE_SECS, index=3
        )  # 20"
        seg_off = st.selectbox(
            "Segundaje en Bloque OFF PRIME:", options=AVAILABLE_SECS, index=5
        )  # 30"
    else:
        seg_unico = st.selectbox(
            "Segundaje único:", options=AVAILABLE_SECS, index=5
        )  # 30"
        seg_prime = seg_unico
        seg_off = seg_unico

    st.markdown("---")
    st.header("📊 Share de Inversión por Canal (SOI)")
    soi_chv = st.slider("% CHV", 0, 100, 30)
    soi_mega = st.slider("% Mega", 0, 100, 30)
    soi_c13 = st.slider("% Canal 13", 0, 100, 25)
    soi_tvn = st.slider("% TVN", 0, 100, 15)
    soi_total = soi_chv + soi_mega + soi_c13 + soi_tvn
    if soi_total != 100:
        st.warning(f"Suma de SOI = {soi_total}% (Ajustar a 100%)")

    soi_dict = {
        "CHV": (soi_chv / 100.0) if soi_total == 100 else 0.30,
        "Mega": (soi_mega / 100.0) if soi_total == 100 else 0.30,
        "Canal 13": (soi_c13 / 100.0) if soi_total == 100 else 0.25,
        "TVN": (soi_tvn / 100.0) if soi_total == 100 else 0.15,
    }

    max_spots_dia_prog = st.number_input(
        "Máx. spots por programa por día:", min_value=1, max_value=5, value=1
    )

# -------------------------------------------------------------
# 3. SELECCIÓN INTELIGENTE DE PROGRAMAS POR CANAL Y BLOQUE
# -------------------------------------------------------------
st.subheader("📋 Inventario de Programas Seleccionados")

inventory_rows = []
for p in progs_db:
    seg_aplicado = seg_prime if p["Bloque"] == "PRIME" else seg_off
    tarifa_val = p["Tarifas"].get(seg_aplicado, 0.0)
    cpp_est = (tarifa_val / p["Rating"]) if p["Rating"] > 0 else 0

    inventory_rows.append(
        {
            "Canal": p["Canal"],
            "Programa": p["Programa"],
            "Bloque": p["Bloque"],
            "Días": p["Días"],
            "Horario": p["Horario"],
            "Segundos": seg_aplicado,
            "Tarifa": tarifa_val,
            "Rating": round(p["Rating"], 2),
            "CPP Est.": round(cpp_est),
        }
    )

df_inv = pd.DataFrame(inventory_rows)

default_prime_set = {
    "Chv Noticias Central", "Primer Plano", "Fiebre de Baile", "Primer Plano Prime",
    "Meganoticias Prime", "Teleserie Prime", "Meganoticias Prime Sabado", "Meganoticias Prime Domingo",
    "Teletrece", "Vecinos al Límite", "Que Dice Chile Prime", "Cultura Prime", "Vertigo",
    "24 Horas Central", "Ahora Caigo Prime", "24 Horas Central Sabado", "24 Horas Central Domingo"
}
default_off_set = {
    "Contigo en La Mañana A", "Contigo en Directo", "Chv Noticias Tarde",
    "Mucho Gusto", "Meganoticias Alerta Semana", "Meganoticias Alerta Finde",
    "Tu Día", "Teletrece Tarde", "Que Dice Chile", "Yo Soy Betty",
    "Buenos Días a Todos", "Ahora Caigo", "Carmen Gloria"
}

def is_default_selected(row):
    prog_lower = row["Programa"].lower()
    if row["Bloque"] == "PRIME":
        return any(k.lower() in prog_lower for k in default_prime_set)
    else:
        return any(k.lower() in prog_lower for k in default_off_set)

df_inv["En Pauta"] = df_inv.apply(is_default_selected, axis=1)

selected_programs_df = st.data_editor(
    df_inv,
    column_config={
        "En Pauta": st.column_config.CheckboxColumn(
            "Incluir", default=False, help="Selecciona los programas para la pauta"
        ),
        "Tarifa": st.column_config.NumberColumn(
            "Valor Unitario (CLP)", format="$%d"
        ),
        "Rating": st.column_config.NumberColumn("Rating", format="%.2f"),
        "CPP Est.": st.column_config.NumberColumn("CPP Est.", format="$%d"),
    },
    disabled=["Canal", "Programa", "Bloque", "Días", "Horario", "Segundos"],
    hide_index=True,
    use_container_width=True,
)

active_progs = selected_programs_df[selected_programs_df["En Pauta"]].copy()

# Limitar a máximo 24 programas para respetar la plantilla física del Excel base
if len(active_progs) > 24:
    balanced_list = []
    for ch in ["CHV", "Mega", "Canal 13", "TVN"]:
        sub_ch = active_progs[active_progs["Canal"] == ch]
        sub_prime = sub_ch[sub_ch["Bloque"] == "PRIME"].sort_values("CPP Est.")
        sub_off = sub_ch[sub_ch["Bloque"] == "OFF PRIME"].sort_values("CPP Est.")
        balanced_list.extend(sub_prime.head(3).to_dict("records"))
        balanced_list.extend(sub_off.head(3).to_dict("records"))
    active_progs = pd.DataFrame(balanced_list).head(24)

# -------------------------------------------------------------
# 4. OPTIMIZADOR CON TOPE DE PRESUPUESTO & DÍAS PERMITIDOS
# -------------------------------------------------------------
dias_list = [start_date + timedelta(days=i) for i in range(num_dias_campana)]
dias_map_es = {0: "L", 1: "M", 2: "W", 3: "J", 4: "V", 5: "S", 6: "D"}

def is_day_allowed(dia_char, dias_tarifa_str):
    d = str(dias_tarifa_str).upper()
    if "L-D" in d or "LMWJVD" in d:
        return True
    if "L-V" in d and dia_char in ["L", "M", "W", "J", "V"]:
        return True
    if "S-D" in d and dia_char in ["S", "D"]:
        return True
    if dia_char in d:
        return True
    return False

def optimize_pauta(progs_df, dates, budget_meta, soi_map, prime_target_pct, max_per_day):
    spots_grid = {p["Programa"]: [0] * len(dates) for _, p in progs_df.iterrows()}
    chan_budget = {ch: budget_meta * pct for ch, pct in soi_map.items()}
    chan_spent = {ch: 0.0 for ch in chan_budget}

    candidates = []
    for _, prog in progs_df.iterrows():
        p_name = prog["Programa"]
        ch = prog["Canal"]
        tariff = float(prog["Tarifa"])
        rating = float(prog["Rating"])
        is_prime = (prog["Bloque"] == "PRIME")
        cpp = (tariff / rating) if rating > 0 else 999999999

        for d_i, dia_dt in enumerate(dates):
            dia_char = dias_map_es[dia_dt.weekday()]
            if is_day_allowed(dia_char, prog["Días"]) and tariff > 0:
                candidates.append({
                    "prog": p_name,
                    "canal": ch,
                    "d_idx": d_i,
                    "tariff": tariff,
                    "rating": rating,
                    "is_prime": is_prime,
                    "cpp": cpp,
                })

    total_spent = 0.0
    total_trps = 0.0
    prime_trps = 0.0

    # Fase 1: Asignar PRIME priorizando menor CPP y respetando hasta 88% del presupuesto del canal
    prime_cands = [c for c in candidates if c["is_prime"]]
    prime_cands.sort(key=lambda x: x["cpp"])

    for cand in prime_cands:
        ch = cand["canal"]
        cost = cand["tariff"]
        p_name = cand["prog"]
        d_idx = cand["d_idx"]

        if spots_grid[p_name][d_idx] >= max_per_day:
            continue
        if chan_spent[ch] + cost > chan_budget[ch] * 0.88:
            continue
        if total_spent + cost > budget_meta * 0.82:
            continue

        spots_grid[p_name][d_idx] += 1
        chan_spent[ch] += cost
        total_spent += cost
        total_trps += cand["rating"]
        prime_trps += cand["rating"]

    # Fase 2: Asignar OFF-PRIME para completar el presupuesto y balancear los canales
    off_cands = [c for c in candidates if not c["is_prime"]]
    off_cands.sort(key=lambda x: x["cpp"])

    for cand in off_cands:
        ch = cand["canal"]
        cost = cand["tariff"]
        p_name = cand["prog"]
        d_idx = cand["d_idx"]

        if spots_grid[p_name][d_idx] >= max_per_day:
            continue
        if chan_spent[ch] + cost > chan_budget[ch] * 1.05:
            continue
        if total_spent + cost > budget_meta:
            continue

        spots_grid[p_name][d_idx] += 1
        chan_spent[ch] += cost
        total_spent += cost
        total_trps += cand["rating"]

    # Fase 3: Relleno fino con cualquier spot disponible hasta alcanzar el presupuesto
    all_sorted = sorted(candidates, key=lambda x: x["cpp"])
    for cand in all_sorted:
        cost = cand["tariff"]
        p_name = cand["prog"]
        d_idx = cand["d_idx"]

        if spots_grid[p_name][d_idx] >= max_per_day:
            continue
        if total_spent + cost > budget_meta:
            continue

        spots_grid[p_name][d_idx] += 1
        chan_spent[cand["canal"]] += cost
        total_spent += cost
        total_trps += cand["rating"]
        if cand["is_prime"]:
            prime_trps += cand["rating"]

    return spots_grid

# Ejecutar optimización ajustada
spots_distrib = optimize_pauta(
    active_progs, dias_list, presupuesto_total, soi_dict, sov_prime_target, max_spots_dia_prog
)

# Totales calculados de la simulación
active_progs["Total_Spots"] = active_progs["Programa"].map(lambda p: sum(spots_distrib.get(p, [])))
active_progs["Inversion"] = active_progs["Total_Spots"] * active_progs["Tarifa"]
active_progs["TRPs"] = active_progs["Total_Spots"] * active_progs["Rating"]

inv_total_sim = active_progs["Inversion"].sum()
trps_totales_sim = active_progs["TRPs"].sum()
cpp_sim = (inv_total_sim / trps_totales_sim) if trps_totales_sim > 0 else 0

trps_prime_sim = active_progs[active_progs["Bloque"] == "PRIME"]["TRPs"].sum()
pct_prime_sim = ((trps_prime_sim / trps_totales_sim) if trps_totales_sim > 0 else 0.0)

# Tarjetas de resumen en pantalla
colA, colB, colC, colD = st.columns(4)
colA.metric(
    "Inversión Proyectada",
    f"${inv_total_sim:,.0f} CLP",
    delta=f"${inv_total_sim - presupuesto_total:,.0f} vs Presupuesto",
    delta_color="normal" if abs(inv_total_sim - presupuesto_total) < 2000000 else "inverse",
)
colB.metric(
    "TRPs Totales",
    f"{trps_totales_sim:.1f} pts",
    delta=f"{trps_totales_sim - objetivo_trps:+.1f} vs Meta",
)
colC.metric("CPP Promedio Ponderado", f"${cpp_sim:,.0f} CLP")
colD.metric(
    "% TRPs en Prime",
    f"{pct_prime_sim*100:.1f}%",
    delta=f"{(pct_prime_sim - sov_prime_target)*100:+.1f}% vs Meta",
)

# -------------------------------------------------------------
# 5. GENERACIÓN DEL EXCEL BASE CON FINES DE SEMANA DINÁMICOS Y FÓRMULAS
# -------------------------------------------------------------
def export_to_excel_base(cliente, campana, version, start_d, end_d, active_p_df, distrib_dict):
    wb = openpyxl.load_workbook("Excel base.xlsx")
    ws = wb["Pauta mes tipo"]

    # Paletas de color
    gray_weekend_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    white_weekday_fill = PatternFill(fill_type=None)
    disabled_day_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")

    # 1. Metadatos
    meses_es = [
        "", "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
        "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
    ]
    nombre_mes = f"{meses_es[start_d.month]} {start_d.year}"

    ws["C10"] = cliente
    ws["C11"] = f"{start_d.strftime('%d/%m/%Y')} al {end_d.strftime('%d/%m/%Y')}"
    ws["C12"] = campana
    ws["C13"] = version
    ws["J15"] = nombre_mes

    dias_semana_abrev = {0: "L", 1: "M", 2: "W", 3: "J", 4: "V", 5: "S", 6: "D"}

    # 2. Configurar Días y formatear dinámicamente Sábado/Domingo en gris
    for i in range(30):
        col_idx = 10 + i  # Columnas J (10) a AM (39)
        if i < len(dias_list):
            cur_dt = dias_list[i]
            day_char = dias_semana_abrev[cur_dt.weekday()]
            is_weekend = cur_dt.weekday() in [5, 6]  # 5=Sábado, 6=Domingo

            ws.cell(16, col_idx, value=day_char)
            ws.cell(17, col_idx, value=cur_dt.day)

            # Aplicar color a toda la columna (filas 16 a 43)
            cell_fill = gray_weekend_fill if is_weekend else white_weekday_fill
            for r in range(16, 44):
                ws.cell(r, col_idx).fill = cell_fill
        else:
            # Días fuera del rango de campaña
            ws.cell(16, col_idx, value=None)
            ws.cell(17, col_idx, value=None)
            for r in range(16, 44):
                ws.cell(r, col_idx).fill = disabled_day_fill

    # 3. Limpiar contenido previo de filas 18 a 42
    for r in range(18, 43):
        for c in range(2, 10):  # B a I
            ws.cell(r, c).value = None
        for c in range(10, 40):  # J a AM
            ws.cell(r, c).value = None
        ws.cell(r, 41).value = None  # AO (Pond)
        ws.cell(r, 42).value = None  # AP (Rating)
        ws.cell(r, 44).value = None  # AR (Valor unitario)
        ws.cell(r, 47).value = None  # AU (Trps prime)

    # 4. Poblar programas con fórmulas activas
    for i, (_, row_p) in enumerate(active_p_df.iterrows()):
        curr_row = 18 + i
        p_name = row_p["Programa"]

        ws.cell(curr_row, 2, value=row_p["Canal"])
        ws.cell(curr_row, 3, value=p_name)
        ws.cell(curr_row, 4, value="Prime" if row_p["Bloque"] == "PRIME" else "Off")
        ws.cell(curr_row, 5, value=row_p["Días"])
        ws.cell(curr_row, 6, value="Spot")
        ws.cell(curr_row, 7, value=row_p["Horario"])
        ws.cell(curr_row, 9, value=int(row_p["Segundos"]))

        # Llenar spots diarios en columnas J a AM
        daily_spots = distrib_dict.get(p_name, [])
        for d_i, num_s in enumerate(daily_spots[:30]):
            if num_s > 0:
                ws.cell(curr_row, 10 + d_i, value=num_s)

        # Fórmulas activas estándar MADCOM
        ws.cell(curr_row, 40, value=f"=SUM(J{curr_row}:AM{curr_row})")
        ws.cell(curr_row, 41, value=1)  # Pond = 1 (Pauta libre)
        ws.cell(curr_row, 42, value=float(row_p["Rating"]))
        ws.cell(curr_row, 43, value=f"=+AP{curr_row}*AN{curr_row}")
        ws.cell(curr_row, 44, value=float(row_p["Tarifa"]))
        ws.cell(curr_row, 45, value=f"=AR{curr_row}*AN{curr_row}")
        ws.cell(curr_row, 46, value=f"=IF(AQ{curr_row}>0, AS{curr_row}/AQ{curr_row}, 0)")
        ws.cell(curr_row, 47, value=f'=IF(D{curr_row}="Prime", AQ{curr_row}, 0)')

    # 5. Fila 43: Totales y TRPS POR DÍA dinámicos
    for i in range(min(len(dias_list), 30)):
        col_letter = openpyxl.utils.get_column_letter(10 + i)
        ws.cell(43, 10 + i, value=f"=SUMPRODUCT({col_letter}18:{col_letter}42,$AP$18:$AP$42)")

    ws.cell(43, 40, value=f"=SUM(AN18:AN42)")
    ws.cell(43, 43, value=f"=SUM(AQ18:AQ42)")
    ws.cell(43, 45, value=f"=SUM(AS18:AS42)")
    ws.cell(43, 46, value=f"=IF(AQ43>0, AS43/AQ43, 0)")
    ws.cell(43, 47, value=f"=SUM(AU18:AU42)")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# -------------------------------------------------------------
# 6. DESCARGA DEL ARCHIVO EXCEL
# -------------------------------------------------------------
st.markdown("---")
st.subheader("📥 Exportación Oficial MADCOM")

excel_buffer = export_to_excel_base(
    cliente_sel,
    campana_nombre,
    version_pauta,
    start_date,
    end_date,
    active_progs,
    spots_distrib,
)

nombre_archivo = f"Pauta_TV_{cliente_sel.replace(' ', '_')}_{start_date.strftime('%d%b')}_{end_date.strftime('%d%b')}.xlsx"

st.download_button(
    label=f"📊 Descargar {nombre_archivo} (Excel Base con Fórmulas)",
    data=excel_buffer,
    file_name=nombre_archivo,
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)
